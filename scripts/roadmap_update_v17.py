"""Roadmap v1.7 (2026-10-26): refleja el cierre de la sincronización LOCAL -> GITHUB.

Actualiza las tres proyecciones en el mismo commit, como exige docs/roadmap/README.md:
  - docs/roadmap/ROADMAP_MASTER_DATA.json  (filas TAP-072/073/075/076 + notas + version)
  - docs/roadmap/ROADMAP_MASTER.csv        (mismas filas)
  - docs/roadmap/ROADMAP_TRADING_AGENTIC_PORTABLE.xlsx (hojas ROADMAP, HISTORIAL, RESUMEN)

Ejecutar desde la raíz del worktree del roadmap:
  xl-venv/Scripts/python.exe scripts_roadmap_v17.py
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl

ROOT = Path.cwd()
JSON_P = ROOT / "docs/roadmap/ROADMAP_MASTER_DATA.json"
CSV_P = ROOT / "docs/roadmap/ROADMAP_MASTER.csv"
XLSX_P = ROOT / "docs/roadmap/ROADMAP_TRADING_AGENTIC_PORTABLE.xlsx"
README_P = ROOT / "docs/roadmap/README.md"

VERSION = "v1.7"
DATE = "2026-10-26"

# Se rellena en update_csv() para mapear las columnas del Excel por posición.
CSV_FIELDS: list[str] = []

UPDATES: dict[str, dict[str, str]] = {
    "TAP-072": {
        "Gate": "0 commits locales sin publicar (git rev-list --branches --not --remotes=origin = 0); "
                "35/35 ramas con SHA local == SHA remoto; respaldo SHA-256 previo a tocar Git",
        "Evidencia": "5b63f265 / LOCAL_GITHUB_SYNC_AUDIT.md + .json + FINAL_GITHUB_SYNC_REPORT.md + .json "
                     "+ EXCLUDED_FROM_SYNC_MANIFEST.json (publicados en GitHub)",
        "Actualizacion": DATE,
        "Siguiente_paso": "Revisar y cerrar los 31 PR draft (#19-#35 y #36-#49)",
        "Notas": "SYNC_COMPLETE != MAIN_INTEGRATED != RUNTIME_CERTIFIED != PROFITABLE. Sin reset --hard, "
                 "clean -fd, push --force ni rebase; ningún worktree, rama, stash o evidencia borrado. "
                 "v1.7: los cuatro reportes de sincronización se publicaron en la rama "
                 "sync/uncommitted-main-20261025 (PR #19); 14 ramas de preservación recibieron PR #36-#49.",
    },
    "TAP-073": {
        "Gate": "ruff check 130 errores y ruff format 444 ficheros en e67194e7; mypy 297 errores; "
                "pytest 79 failed / 3274 passed / 34 skipped / 14 errors",
        "Evidencia": "checkout limpio gate-main-20261026 de e67194e7 con PYTHONPATH al src limpio, "
                     "2026-10-26; ruff/format/mypy idénticos a los de v1.6 (2026-10-25, 90 failed / 3263 passed)",
        "Actualizacion": DATE,
        "Notas": "Falla preexistente: NO la introduce la sincronización. Re-medición 2026-10-26: mismos "
                 "130/444/297 que en v1.6 y 19 de 20 ficheros con fallos coinciden con la baseline "
                 "(tests/unit/research/**). La CI Baseline ya fallaba sobre origin/main desde 2026-07-10. "
                 "origin/main 5e6f9632 sí está en verde (582 tests PASS) pero está 263 commits por detrás.",
    },
    "TAP-075": {
        "Gate": "31 PR draft abiertos (#19-#35 y #36-#49); cabeceras == SHA de rama local",
        "Evidencia": "refs/pull/*/head verificados con git ls-remote y API /pulls?state=open (38 PR abiertos)",
        "Actualizacion": DATE,
        "Notas": "Todos en draft: gates NO ejecutados en esas ramas -> NOT_VERIFIED. "
                 "audit/economic-blocker-001 no recibe PR porque es idéntica a main; "
                 "main se integra a través del PR #18.",
    },
    "TAP-076": {
        "Gate": "8 commits en rama sync/uncommitted-main-20261025; árbol limpio; PR #19 draft; "
                "colección rota re-verificada 2026-10-26 (3532 tests + 1 error de import)",
        "Evidencia": "5b63f265; PR #19; respaldo sync-backup-20261025 con SHA256SUMS.txt",
        "Actualizacion": DATE,
        "Notas": "El shim de settings.py (346 -> 27 lineas) rompe la recolección: tests/unit/config/"
                 "test_failfast.py no puede importar YamlDirectorySource (commit c16c241b). "
                 "Publicado, NO integrable. La colección de main e67194e7 sí es correcta (3401 tests, 0 errores).",
    },
}

NOTES_ADD = (
    f"\nRevisión {VERSION} ({DATE}) — cierre de la sincronización LOCAL -> GITHUB: 0 commits locales sin "
    "publicar (35/35 ramas con SHA local == SHA remoto), 31 PR draft (#19-#35 y #36-#49) y los cuatro "
    "reportes (LOCAL_GITHUB_SYNC_AUDIT.md/.json + FINAL_GITHUB_SYNC_REPORT.md/.json) publicados. "
    "Gates re-meditos en checkout limpio de e67194e7: ruff 130 / format 444 / mypy 297 (idénticos a v1.6) y "
    "pytest 79 failed / 3274 passed / 34 skipped / 14 errors, con 19 de 20 ficheros coincidentes con la "
    "baseline: el rojo es preexistente. Sigue BLOQUEADO el PR #19 por error de colección y sigue sin "
    "integrarse origin/main (263 commits por detrás). Nada pasa a CERTIFIED."
)

HISTORIAL_ROW = [
    DATE,
    f"{VERSION} — sincronización LOCAL -> GITHUB cerrada",
    "35/35 ramas publicadas y verificadas, 0 commits sin publicar, 31 PR draft (#19-#49), "
    "4 reportes publicados",
    "TAP-072, TAP-073, TAP-075, TAP-076",
    "Gates re-meditos: ruff 130 / format 444 / mypy 297 / pytest 79F-3274P-14E (preexistente)",
    "Codebuff",
    "SYNC_COMPLETE != MAIN_INTEGRATED != RUNTIME_CERTIFIED != PROFITABLE",
]


def update_json() -> None:
    data = json.loads(JSON_P.read_text(encoding="utf-8"))
    data["version"] = VERSION
    data["updated"] = DATE
    touched = set()
    for row in data["rows"]:
        upd = UPDATES.get(row["ID"])
        if not upd:
            continue
        for field, value in upd.items():
            if field in row:
                row[field] = value
                touched.add(f"{row['ID']}.{field}")
    if NOTES_ADD.strip() not in data["notes"]:
        data["notes"] = data["notes"].rstrip() + NOTES_ADD
    JSON_P.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"json: {len(data['rows'])} filas, actualizadas {len(touched)} celdas, version {data['version']}")


def update_csv() -> None:
    with CSV_P.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
        fields = list(rows[0].keys())
    global CSV_FIELDS
    CSV_FIELDS = fields
    touched = 0
    for row in rows:
        upd = UPDATES.get(row["ID"])
        if not upd:
            continue
        for field, value in upd.items():
            if field in row:
                row[field] = value
                touched += 1
    with CSV_P.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, lineterminator="\r\n")
        writer.writeheader()
        writer.writerows(rows)
    print(f"csv: {len(rows)} filas, actualizadas {touched} celdas")


def _header_row(ws) -> int:
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=12, max_col=1), start=1):
        if row[0].value == "ID":
            return idx
    raise RuntimeError(f"no se encontró la cabecera 'ID' en {ws.title}")


def update_xlsx() -> None:
    wb = openpyxl.load_workbook(XLSX_P)

    # El orden de columnas del Excel es el mismo que el del CSV (ver README), pero los
    # rótulos difieren (p. ej. "Frente / línea"), así que se mapea por posición.
    fields = CSV_FIELDS or [c.value for c in openpyxl.load_workbook(XLSX_P)["ROADMAP"][5] if c.value]

    ws = wb["ROADMAP"]
    header_idx = _header_row(ws)
    touched = 0
    for row in ws.iter_rows(min_row=header_idx + 1):
        upd = UPDATES.get(row[0].value)
        if not upd:
            continue
        for field, value in upd.items():
            if field in fields:
                ws.cell(row=row[0].row, column=fields.index(field) + 1, value=value)
                touched += 1
    # Cabecera de la hoja (fecha de corte + versión)
    ws["A3"] = f"{DATE} · {VERSION} · sincronización LOCAL -> GITHUB cerrada: 35/35 ramas publicadas, " \
               "31 PR draft, gates en rojo preexistente"

    hist = wb["HISTORIAL"]
    already = any(
        str(r[0].value) == DATE for r in hist.iter_rows(min_row=6, max_col=1) if r[0].value
    )
    if not already:
        hist.append(HISTORIAL_ROW)

    resumen = wb["RESUMEN"]
    resumen["A3"] = f"Corte GitHub: {DATE} · {VERSION} · 0 commits sin publicar · 31 PR draft · " \
                    "origin/main 263 commits por detrás"

    wb.save(XLSX_P)
    print(f"xlsx: {touched} celdas actualizadas, HISTORIAL {'ya existía' if already else '+1 fila'}, "
          "RESUMEN y cabecera actualizados")


def update_readme() -> None:
    text = README_P.read_text(encoding="utf-8")
    text = text.replace(
        "**Fecha de corte:** 2026-10-25 · **Versión actual:** 1.6",
        f"**Fecha de corte:** {DATE} · **Versión actual:** {VERSION.replace('v', '')}",
        1,
    )
    marker = "## Estado destacado del corte inicial"
    add = ""
    if f"### {VERSION}" not in text:
        add = (
        f"### v1.7 ({DATE}) — cierre de la sincronización LOCAL → GITHUB\n\n"
        "- 35/35 ramas locales publicadas y verificadas (`SHA local == SHA remoto`), 0 commits sin publicar, "
        "15 ramas de preservación con PR `#19` y `#36`-`#49`.\n"
        "- Los cuatro reportes (`LOCAL_GITHUB_SYNC_AUDIT.md/.json`, `FINAL_GITHUB_SYNC_REPORT.md/.json`) y el "
        "manifiesto de exclusiones quedan versionados en `sync/uncommitted-main-20261025`.\n"
        "- Gates re-meditos en checkout limpio de `e67194e7`: ruff 130 / format 444 / mypy 297 (idénticos a v1.6) "
        "y pytest 79 failed / 3274 passed / 34 skipped / 14 errors — rojo **preexistente**, no introducido por "
        "la sincronización.\n"
        "- Sigue `BLOQUEADO` el PR `#19` (error de colección en `test_failfast.py`); `origin/main` sigue 263 "
        "commits por detrás.\n\n"
    )
    if add and marker in text:
        text = text.replace(marker, add + marker, 1)
    elif add:
        text = text.rstrip() + "\n\n" + add
    README_P.write_text(text, encoding="utf-8")
    print(f"readme: {'changelog añadido' if add else 'changelog ya presente'}, cabecera {VERSION}")


def main() -> None:
    update_json()
    update_csv()
    update_xlsx()
    update_readme()


if __name__ == "__main__":
    main()
